"""감독관: 덱을 관문마다 살피고, 고칠 수 있는 것은 고침을 내놓고, 위험하면 굽기를 멈춘다.

수준
- 멈춤: 이대로 내보내면 안 됨(넘침, 견본 글, 숫자의 정직, 없는 그림 …). 굽기 전에 묻는다
- 주의: 고치는 게 좋음(글자 수 한도, 대비, 해상도 …)
- 알림: 알아 두면 좋음(같은 그림 되풀이, 발표 시간 …)

찾은 것은 {"규칙", "수준", "장"(1부터, 덱 전체면 0), "말", "고침"(있으면 고치는 방법 이름)}.
사람이 '괜찮아요'를 누른 것은 설계의 "감독관무시"에 남아 다음부터 안 보인다.

시험: python 감독관.py 시험   (규칙마다 틀린 예·맞는 예 + 일부러 틀린 설계 5개, 늘 0 틀림이어야 한다)
"""

from __future__ import annotations

import json
import re
import sys
from pathlib import Path

import 배치
import 소재
import 스타일 as 스

수준차례 = {"멈춤": 0, "주의": 1, "알림": 2}
견본표시 = re.compile(r"넣으세요")
_그림칸 = ("그림", "그림여럿")


# ---------------------------------------------------------------- 도움
def _값글(v) -> list[str]:
    if v is None:
        return []
    if isinstance(v, list):
        return [str(x) for x in v if str(x).strip()]
    return [str(v)] if str(v).strip() else []


def _대비(a: str, b: str) -> float:
    la, lb = sorted((스._밝기(a), 스._밝기(b)), reverse=True)
    return (la + 0.05) / (lb + 0.05)


def _섞기(아래: str, 위: str, a: float) -> str:
    x = [int(아래[i:i + 2], 16) for i in (1, 3, 5)]
    y = [int(위[i:i + 2], 16) for i in (1, 3, 5)]
    return "#" + "".join(f"{round(p + (q - p) * a):02X}" for p, q in zip(x, y))


def _글실제(b: dict) -> tuple[float, float, float, float]:
    """글 박스에서 글이 실제로 차지하는 자리(세로 맞춤 반영)."""
    h = min(b["h"], 배치.글높이(b["글"], b["w"], b["크기"], b["굵기"], b["점"]))
    y = b["y"] + {"위": 0, "가운데": (b["h"] - h) / 2, "아래": b["h"] - h}[b["세로"]]
    return b["x"], y, b["w"], h


def _안에(px: float, py: float, b: dict) -> bool:
    return b["x"] <= px <= b["x"] + b["w"] and b["y"] <= py <= b["y"] + b["h"]


def _글바탕(박스들: list[dict], i: int, 색: dict) -> str | None:
    """i번째 글 박스 아래의 색. 사진 위(막 없이)면 None(잴 수 없음)."""
    b = 박스들[i]
    cx, cy = b["x"] + b["w"] / 2, b["y"] + b["h"] / 2
    지금: str | None = 색["바탕"]
    for a in 박스들[:i]:
        if not _안에(cx, cy, a):
            continue
        if a["종류"] == "그림":
            지금 = None
        elif a["종류"] == "네모" and a["채움"] in 색 and a.get("모양") != "화살":
            투명 = a.get("투명") or 0
            지금 = 색[a["채움"]] if not 투명 else _섞기(지금 or "#808080", 색[a["채움"]], 1 - 투명)
    return 지금


_v1기억: dict = {}


def _v1장수(f: Path) -> int | None:
    """v1 PPTX 장 수(저장할 때마다 살피므로 파일 시각으로 기억해 둔다)."""
    try:
        열쇠 = (str(f), f.stat().st_mtime)
    except OSError:
        return None
    if 열쇠 not in _v1기억:
        from pptx import Presentation
        _v1기억[열쇠] = len(Presentation(str(f)).slides)
    return _v1기억[열쇠]


# ---------------------------------------------------------------- 규칙
def 살피기(설계: dict, 덱: dict, 프로젝트: Path | None = None) -> list[dict]:
    """설계와 배치된 덱(엔진.덱배치 결과)을 살핀다."""
    찾음: list[dict] = []

    def 더하기(규칙, 수준, 장, 말, 고침=None):
        찾음.append({"규칙": 규칙, "수준": 수준, "장": 장, "말": 말, **({"고침": 고침} if 고침 else {})})

    장들 = 설계.get("장", [])
    if not 장들:
        더하기("빈덱", "멈춤", 0, "장이 하나도 없어요")
    # 발표 시간: 한 장에 30초~2분
    분 = 설계.get("발표분")
    if 분 and 장들:
        if len(장들) > 2 * 분:
            더하기("발표시간", "주의", 0, f"{분}분 발표에 {len(장들)}장은 많아요(1분에 2장 넘음)")
        elif len(장들) < 분 / 3:
            더하기("발표시간", "알림", 0, f"{분}분 발표에 {len(장들)}장이면 한 장을 3분 넘게 봐요")

    그림쓰임: dict[str, list[int]] = {}
    쓴원본: dict = {}
    v1장수 = _v1장수(프로젝트 / 설계["바탕"]["파일"]) if 프로젝트 is not None and 설계.get("바탕") else None
    본것: dict[str, int] = {}
    for n, (장, 배치된) in enumerate(zip(장들, 덱["장들"]), start=1):
        종류 = 장.get("종류")
        틀 = 배치.장종류.get(종류)
        if 종류 == "원본장":            # v1 장: 있는 번호인지, 두 번 쓰지 않았는지
            원본 = 장.get("원본")
            if v1장수 is not None and not (isinstance(원본, int) and 1 <= 원본 <= v1장수):
                더하기("원본장", "멈춤", n, f"원본 {원본}장은 v1에 없어요(v1은 {v1장수}장)", "장빼기")
            elif 원본 in 쓴원본:
                더하기("원본장", "멈춤", n, f"원본 {원본}장을 {쓴원본[원본]}장에서 이미 썼어요. 한 장은 한 번만 넣을 수 있어요", "장빼기")
            else:
                쓴원본[원본] = n
            continue
        if 틀 is None:
            더하기("모르는종류", "주의", n, f"'{종류}'는 모르는 장 종류라서 글그림으로 그렸어요")
            틀 = 배치.장종류["글그림"]
        # --- 설계: 필수 칸, 견본 글, 글자 수
        for 칸 in 틀["칸"]:
            v = 장.get(칸["이름"])
            if 칸.get("필수") and not (_값글(v) if 칸["형태"] not in _그림칸 else v):
                더하기("필수칸", "멈춤", n, f"'{칸['이름']}' 칸이 비었어요")
            if 칸["형태"] in ("한줄", "여러줄"):
                줄들 = _값글(v)
                if any(견본표시.search(x) for x in 줄들):
                    더하기("견본글", "멈춤", n, f"'{칸['이름']}'에 견본 글이 남았어요")
                한도 = 칸.get("한도")
                긴줄 = [x for x in 줄들 if 한도 and len(x) > 한도]
                if 긴줄:
                    더하기("글자수", "주의", n, f"'{칸['이름']}'이 {len(긴줄[0])}자예요(한도 {한도}자). 줄이면 더 잘 읽혀요")
                if 칸["형태"] == "여러줄" and 칸.get("줄수") and len(줄들) > 칸["줄수"]:
                    고침 = "장나누기" if 칸["이름"] == "글" and 종류 in ("글만", "글그림") else None
                    더하기("줄수", "주의", n, f"'{칸['이름']}'이 {len(줄들)}줄이에요(한도 {칸['줄수']}줄)", 고침)
        # --- 숫자의 정직
        if 종류 == "도표":
            import 엔진          # 표 파일을 붙이는 일은 엔진이 한다(서로 부르지 않게 여기서)
            d = 배치.도표자료(엔진._표붙이기(장, 프로젝트, 설계))
            if not d or not d.get("계열"):
                더하기("도표자료", "멈춤", n, "도표에 그릴 숫자 열이 없어요. 표를 골라 주세요")
            elif d["모양"] == "막대" and 장.get("축0") is False:
                더하기("축0", "멈춤", n, "막대 도표는 0에서 시작해야 길이가 값과 비례해요", "축0켜기")
            if d and 장.get("합100"):
                for i in range(len(d.get("가로", []))):
                    합 = sum(s["값"][i] or 0 for s in d["계열"])
                    if abs(합 - 100) > 1:
                        더하기("합100", "멈춤", n, f"'{d['가로'][i]}' 줄의 합이 {합:g}예요(100이어야 해요)")
                        break
        if 종류 in ("숫자", "결과해석"):
            for 값, _ in 배치._숫자들(장.get("숫자들")):
                if 값 and not re.search(r"\d", 값):
                    더하기("숫자없음", "주의", n, f"'{값}'에 숫자가 없어요. 숫자 장은 숫자로 말해요")
        # --- 그림: 있는지, 해상도
        for 칸 in 틀["칸"]:
            if 칸["형태"] not in _그림칸:
                continue
            경로들 = 장.get(칸["이름"]) or []
            for 경로 in ([경로들] if isinstance(경로들, str) else 경로들):
                if not isinstance(경로, str) or not 경로.strip():
                    continue                     # 빈 칸은 '필수칸'이 따로 잡는다
                그림쓰임.setdefault(경로, [])
                if n not in 그림쓰임[경로]:
                    그림쓰임[경로].append(n)
                if 프로젝트 is not None and not (프로젝트 / 경로).is_file():
                    고침 = "글만으로" if 종류 == "글그림" else None
                    더하기("그림없음", "멈춤", n, f"그림 파일이 없어요: {경로}", 고침)
        for b in 배치된["박스"]:
            if b["종류"] == "그림" and 프로젝트 is not None and (프로젝트 / b["그림"]).is_file():
                크기 = 소재.크기안전(프로젝트 / b["그림"])
                if 크기 is None:
                    더하기("그림읽기", "멈춤", n, f"{Path(b['그림']).name}을 그림으로 읽지 못했어요"
                        "(깨진 파일이거나 PPTX에 못 넣는 형식). PNG·JPG·TIFF로 바꿔 넣어 주세요")
                    continue
                iw, ih = 크기
                if b["맞춤"] == "채우기":       # 잘리는 비율: 가장자리 스케일바·정보 띠가 사라질 수 있다
                    잘림 = 1 - min((b["w"] / b["h"]) / (iw / ih), (iw / ih) / (b["w"] / b["h"]))
                    if 잘림 > 0.03:
                        더하기("잘림", "주의", n, f"{Path(b['그림']).name}의 가장자리 {잘림:.0%}가 잘려요"
                            "(SEM 스케일바가 사라질 수 있어요)", "담기로")
                k = (min if b["맞춤"] == "담기" else max)(b["w"] / iw, b["h"] / ih)
                if k > 1.5:
                    더하기("해상도", "주의", n, f"{Path(b['그림']).name}({iw}×{ih})을 {k:.1f}배로 키워 흐려 보여요")
        for b in 배치된["박스"]:           # 표: 가장 작은 글자로도 칸·줄이 안 들어가면 글이 잘린다
            if b["종류"] == "표" and b.get("넘침"):
                더하기("표넘침", "멈춤", n, "표가 장에 다 안 들어가요(칸 글이 잘리거나 줄이 너무 많아요). "
                    "칸 글을 줄이거나 표를 두 장으로 나눠 주세요")
        # --- 장 짜기: 넘침, 작은 글자, 대비, 겹침
        글박스 = [(i, b) for i, b in enumerate(배치된["박스"]) if b["종류"] == "글" and b["글"]]
        for i, b in 글박스:
            if b["넘침"]:
                더하기("넘침", "멈춤", n, f"'{b['자리']}' 글이 상자보다 길어요. 줄이거나 장을 나눠 주세요",
                    "장나누기" if b["자리"] == "글" and 종류 in ("글만", "글그림") else None)
            if b["크기"] < 스.글자["가장작은"] - 0.5:
                더하기("작은글자", "주의", n, f"'{b['자리']}' 글자가 {b['크기']:.0f}px예요(24px 넘게)")
            아래 = _글바탕(배치된["박스"], i, 덱["색"])
            if 아래 and b["색"] in 덱["색"]:
                c = _대비(덱["색"][b["색"]], 아래)
                기준 = 3 if b["크기"] >= 48 or (b["크기"] >= 37 and b["굵기"] >= 700) else 4.5
                if c < 기준:
                    더하기("대비", "멈춤" if c < 2 else "주의", n,
                        f"'{b['자리'] or '글'}' 글과 바탕의 대비가 {c:.1f}:1이에요(≥{기준:g}:1)")
        for a in range(len(글박스)):
            for c in range(a + 1, len(글박스)):
                x1, y1, w1, h1 = _글실제(글박스[a][1])
                x2, y2, w2, h2 = _글실제(글박스[c][1])
                겹 = max(0, min(x1 + w1, x2 + w2) - max(x1, x2)) * max(0, min(y1 + h1, y2 + h2) - max(y1, y2))
                if 겹 > 0.05 * min(w1 * h1, w2 * h2):
                    더하기("겹침", "주의", n, f"'{글박스[a][1]['자리']}'와 '{글박스[c][1]['자리']}' 글이 겹쳐요")
        # --- 완성: 빈 장, 되풀이 장
        내용 = {k: v for k, v in 장.items() if k not in ("노트", "시간") and not k.startswith("_")}
        if not any(_값글(v) for k, v in 내용.items() if k != "종류"):
            더하기("빈장", "주의", n, "내용이 없는 장이에요")
        열쇠 = json.dumps(내용, ensure_ascii=False, sort_keys=True)
        if 열쇠 in 본것:
            더하기("되풀이장", "주의", n, f"{본것[열쇠]}장과 똑같아요", "장빼기")
        else:
            본것[열쇠] = n
    for 경로, 쓴장 in 그림쓰임.items():
        if len(쓴장) > 2:
            더하기("그림되풀이", "알림", 쓴장[-1], f"{Path(경로).name} 그림을 {', '.join(map(str, 쓴장))}장에 썼어요")
    # 사람이 괜찮다고 한 것 빼기
    # 말까지 같은 줄만 숨긴다(예전 기록은 말이 없어 규칙·장으로)
    무시 = 설계.get("감독관무시", [])
    def 숨김(f):
        return any(x.get("규칙") == f["규칙"] and x.get("장") == f["장"]
                   and x.get("말", f["말"]) == f["말"] for x in 무시)
    찾음 = [f for f in 찾음 if not 숨김(f)]
    return sorted(찾음, key=lambda f: (수준차례[f["수준"]], f["장"]))


def 멈춤(찾음: list[dict]) -> list[dict]:
    return [f for f in 찾음 if f["수준"] == "멈춤"]


# ---------------------------------------------------------------- 고침 (사람이 눌러야 한다)
def 고치기(설계: dict, 규칙: str, 장번호: int, 고침: str) -> tuple[dict, str]:
    """설계 사본을 고쳐 돌려준다. 지우지 않는다: 뺀 장은 휴지통으로."""
    설계 = json.loads(json.dumps(설계, ensure_ascii=False))
    장들 = 설계.get("장", [])
    i = 장번호 - 1
    if not 0 <= i < len(장들):
        raise ValueError("그 장이 없어요")
    장 = 장들[i]
    if 고침 == "축0켜기":
        장["축0"] = True
        return 설계, f"{장번호}장 도표를 0에서 시작하게 했어요"
    if 고침 == "글만으로":
        (설계.setdefault("휴지통", [])).append(json.loads(json.dumps(장, ensure_ascii=False)))
        장["종류"] = "글만"
        장.pop("그림", None)
        장.pop("그림설명", None)
        return 설계, f"{장번호}장을 그림 없는 글만 장으로 바꿨어요(원래 장은 휴지통에)"
    if 고침 == "장나누기":
        글 = _값글(장.get("글"))
        if len(글) < 2:
            raise ValueError("나눌 줄이 모자라요")
        반 = (len(글) + 1) // 2
        뒤 = {**json.loads(json.dumps(장, ensure_ascii=False)), "글": 글[반:]}
        뒤["제목"] = f"{장.get('제목', '')} (이어서)".strip()
        if 뒤.get("종류") == "글그림":        # 그림은 앞 장에만, 뒤 장은 글만
            뒤["종류"] = "글만"
            뒤.pop("그림", None)
            뒤.pop("그림설명", None)
        장["글"] = 글[:반]
        장들.insert(i + 1, 뒤)
        return 설계, f"{장번호}장을 두 장으로 나눴어요"
    if 고침 == "담기로":
        장["맞춤"] = "담기"
        return 설계, f"{장번호}장 사진을 자르지 않고 전체가 보이게 했어요"
    if 고침 == "장빼기":
        설계.setdefault("휴지통", []).append(장들.pop(i))
        return 설계, f"{장번호}장을 휴지통으로 옮겼어요"
    raise ValueError(f"모르는 고침: {고침}")


def 무시하기(설계: dict, 규칙: str, 장번호: int, 말: str | None = None) -> dict:
    설계 = dict(설계)
    기록 = {"규칙": 규칙, "장": 장번호, **({"말": 말} if 말 else {})}
    설계["감독관무시"] = [*설계.get("감독관무시", []), 기록]
    return 설계


# ---------------------------------------------------------------- 완성 점검 (굽고 나서)
def 완성점검(설계: dict, 결과폴더: Path, 덱: dict) -> list[dict]:
    찾음 = []
    n = len(설계.get("장", []))
    for f in 결과폴더.glob("*.pptx"):
        from pptx import Presentation
        prs = Presentation(str(f))
        if len(prs.slides) != n:
            찾음.append({"규칙": "장수", "수준": "멈춤", "장": 0, "말": f"{f.name}가 {len(prs.slides)}장이에요(설계 {n}장)"})
        elif not 설계.get("바탕"):
            for k, (s, 장) in enumerate(zip(prs.slides, 덱["장들"]), start=1):
                글수 = sum(1 for b in 장["박스"] if b["종류"] == "글" and b["글"])
                살아있는 = sum(1 for sh in s.shapes if sh.has_text_frame and sh.text_frame.text.strip())
                if 살아있는 < 글수:
                    찾음.append({"규칙": "글상자", "수준": "주의", "장": k,
                               "말": f"PPTX에 글상자가 {살아있는}개뿐이에요(화면 {글수}개)"})
    for f in 결과폴더.glob("*.pdf"):
        import pypdfium2 as pdfium
        문서 = pdfium.PdfDocument(str(f))
        쪽 = len(문서)
        문서.close()
        if 쪽 != n:
            찾음.append({"규칙": "장수", "수준": "멈춤", "장": 0, "말": f"{f.name}가 {쪽}쪽이에요(설계 {n}장)"})
    그림들 = list((결과폴더 / "그림").glob("*.png")) if (결과폴더 / "그림").exists() else []
    if 그림들 and len(그림들) != n:
        찾음.append({"규칙": "장수", "수준": "멈춤", "장": 0, "말": f"PNG가 {len(그림들)}장이에요(설계 {n}장)"})
    return 찾음


# ---------------------------------------------------------------- 시험
_시험폴더 = Path(__file__).parent / "감독관시험"


def _배치(설계: dict, 프로젝트: Path | None) -> dict:
    import 엔진
    return 엔진.덱배치(설계, 프로젝트)


def 시험() -> int:
    """규칙 예(감독관시험/규칙예.json) + 틀린 설계(감독관시험/틀린설계_*.json). 틀린 수를 돌려준다."""
    틀림 = 0
    예들 = json.loads((_시험폴더 / "규칙예.json").read_text("utf-8"))
    for 예 in 예들:
        설계 = {"이름": "시험", "장": 예["장"], **예.get("덱", {})}
        찾음 = 살피기(설계, _배치(설계, _시험폴더), _시험폴더)
        규칙들 = {f["규칙"] for f in 찾음}
        if 예["잡아야"] and 예["규칙"] not in 규칙들:
            print(f"✗ 못 잡음: {예['이름']} ({예['규칙']})")
            틀림 += 1
        elif not 예["잡아야"] and 규칙들:
            print(f"✗ 괜히 잡음: {예['이름']} → {sorted(규칙들)}")
            틀림 += 1
        else:
            print(f"✓ {예['이름']}")
    for f in sorted(_시험폴더.glob("틀린설계_*.json")):
        설계 = json.loads(f.read_text("utf-8"))
        찾음 = 살피기(설계, _배치(설계, _시험폴더), _시험폴더)
        빠짐 = set(설계["_잡아야"]) - {x["규칙"] for x in 찾음}
        if 빠짐 or not 멈춤(찾음):
            print(f"✗ {f.name}: 못 잡은 것 {sorted(빠짐)}{'' if 멈춤(찾음) else ', 멈추지 않음'}")
            틀림 += 1
        else:
            print(f"✓ {f.name}: {len(찾음)}개 찾음, 멈춤 {len(멈춤(찾음))}개")
    print(f"틀림 {틀림}")
    return 틀림


if __name__ == "__main__":
    if sys.argv[1:2] == ["시험"]:
        sys.exit(1 if 시험() else 0)
    print(__doc__)
