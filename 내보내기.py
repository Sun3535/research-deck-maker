"""내보내기: 구운 장 그림으로 만드는 나머지 결과들.

- 그림묶음.zip : 장마다 1920×1080 PNG
- 발표영상.mp4 : 장 그림 + 넘김(0.5초 겹쳐 사라지기). 한 장 시간은 장의 "시간"(초), 없으면 덱의 "장시간", 그것도 없으면 5초
- 대본.md      : 장마다 발표자 노트(없으면 화면 글)
- 데이터.xlsx  : 도표·표 장에 쓴 표를 장마다 시트로
- 소재/        : 덱에 쓴 그림 원본(화질 그대로 복사)
ffmpeg는 PC에 있는 것(또는 앱의 도구/ 폴더)을 쓴다. 없으면 영상만 건너뛰고 알린다.
"""

from __future__ import annotations

import shutil
import subprocess
import zipfile
from pathlib import Path

import 재료받기

기본장시간 = 5.0
넘김 = 0.5


def 그림묶음(그림들: list[Path], 저장: Path) -> Path:
    with zipfile.ZipFile(저장, "w", zipfile.ZIP_STORED) as z:     # PNG는 이미 눌려 있어 그대로 담는다
        for f in 그림들:
            z.write(f, f.name)
    return 저장


def 장시간들(설계: dict) -> list[float]:
    기본 = float(설계.get("장시간") or 기본장시간)
    out = []
    for 장 in 설계.get("장", []):
        try:
            out.append(max(1.0, float(장.get("시간") or 기본)))
        except (TypeError, ValueError):
            out.append(기본)
    return out


def 영상(그림들: list[Path], 시간들: list[float], 저장: Path) -> str | None:
    """돌려주는 것: 알릴 말(못 만들었을 때) 또는 None."""
    ff = 재료받기.ffmpeg()
    if not ff:
        return "ffmpeg가 없어 영상은 건너뛰었어요(앱의 도구 폴더에 ffmpeg를 두면 돼요)"
    if not 그림들:
        return "영상으로 만들 장 그림이 없어요"
    입력 = []
    for f, d in zip(그림들, 시간들):
        입력 += ["-loop", "1", "-framerate", "30", "-t", f"{d + (넘김 if len(그림들) > 1 else 0):.2f}", "-i", str(f)]
    if len(그림들) == 1:
        거름 = "[0:v]format=yuv420p[v]"
    else:
        조각, 앞, 지난 = [], "[0:v]", 0.0
        for k in range(1, len(그림들)):
            지난 += 시간들[k - 1]
            뒤 = f"[x{k}]" if k < len(그림들) - 1 else "[xv]"
            조각.append(f"{앞}[{k}:v]xfade=transition=fade:duration={넘김}:offset={지난:.2f}{뒤}")
            앞 = 뒤
        거름 = ";".join(조각) + ";[xv]format=yuv420p[v]"
    명령 = [ff, "-hide_banner", "-loglevel", "error", "-y", *입력, "-filter_complex", 거름, "-map", "[v]",
          "-r", "30", "-c:v", "libx264", "-preset", "medium", "-crf", "18", "-movflags", "+faststart", str(저장)]
    try:
        subprocess.run(명령, check=True, timeout=1800, capture_output=True)
    except (subprocess.CalledProcessError, subprocess.TimeoutExpired) as e:
        말 = getattr(e, "stderr", b"") or b""
        return "영상을 만들지 못했어요: " + 말.decode("utf-8", "replace")[-300:]
    return None


def _화면글(장: dict) -> str:
    조각 = []
    for k, v in 장.items():
        if k in ("종류", "노트", "시간", "그림", "그림들", "왼그림", "오른그림", "표", "모양", "가로", "세로",
                 "축0", "합100", "단위", "맞춤", "원본", "고침") or k.startswith("_"):
            continue
        if isinstance(v, list):
            조각 += [str(x) for x in v if str(x).strip()]
        elif isinstance(v, str) and v.strip():
            조각.append(v)
    return " / ".join(조각)


def 대본(설계: dict, 저장: Path) -> Path:
    시간들 = 장시간들(설계)
    줄 = [f"# {설계.get('이름', '발표')} 대본", "",
          f"장 {len(시간들)}개 · 영상 기준 약 {sum(시간들) / 60:.1f}분", ""]
    for n, (장, t) in enumerate(zip(설계.get("장", []), 시간들), start=1):
        제목 = 장.get("제목") or 장.get("큰말") or 장.get("인용") or 장.get("종류", "")
        if 장.get("종류") == "원본장":
            제목 = f"원본 {장.get('원본')}장"
        줄 += [f"## {n}장 · {제목}", f"*{장.get('종류', '')} · {t:g}초*", ""]
        if 장.get("노트"):
            줄 += [장["노트"], ""]
        else:
            줄 += ["(발표자 노트가 없어요. 화면 글: " + (_화면글(장) or "없음") + ")", ""]
    저장.write_text("\n".join(줄), "utf-8")
    return 저장


def 데이터(설계: dict, 표붙이기, 저장: Path) -> Path | None:
    """도표·표 장의 표 → 장마다 시트. 표가 하나도 없으면 None."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    wb = Workbook()
    목록 = wb.active
    목록.title = "장 목록"
    목록.append(["장", "종류", "제목", "표 파일"])
    있음 = False
    for n, 장 in enumerate(설계.get("장", []), start=1):
        목록.append([n, 장.get("종류"), 장.get("제목") or 장.get("큰말") or 장.get("인용") or "", 장.get("표") or ""])
        if not 장.get("표"):
            continue
        t = 표붙이기(장).get("_표")
        if not t:
            continue
        있음 = True
        이름 = f"{n}장 {장.get('제목', '')}"
        for c in '[]:*?/\\':
            이름 = 이름.replace(c, " ")
        ws = wb.create_sheet(이름[:31])
        ws.append(t["머리"])
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in t["줄들"]:
            ws.append(r)
    if not 있음:
        return None
    wb.save(저장)
    return 저장


def 소재모으기(설계: dict, 프로젝트: Path, 폴더: Path) -> list[Path]:
    """덱에 쓴 그림 원본을 그대로 복사(파일 바이트 그대로 = 무손실)."""
    쓴것: list[str] = []
    for 장 in 설계.get("장", []):
        for k, v in 장.items():
            if k.startswith("_"):
                continue
            for 경로 in ([v] if isinstance(v, str) else v if isinstance(v, list) else []):
                if (isinstance(경로, str) and Path(경로).suffix.lower() in
                        (".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif", ".webp")
                        and 경로 not in 쓴것 and (프로젝트 / 경로).is_file()):
                    쓴것.append(경로)
    out = []
    for 경로 in 쓴것:
        원래 = Path(경로)
        이름, n = 원래.name, 2
        if (폴더 / 이름).exists():           # 이름이 같으면 재료 폴더 이름을, 그래도 같으면 번호를 붙인다
            이름 = f"{원래.parent.parent.name}_{원래.name}" if 원래.parent.name == "그림" else f"{원래.parent.name}_{원래.name}"
        while (폴더 / 이름).exists():
            이름, n = f"{원래.stem}_{n}{원래.suffix}", n + 1
        폴더.mkdir(parents=True, exist_ok=True)
        shutil.copy2(프로젝트 / 경로, 폴더 / 이름)
        out.append(폴더 / 이름)
    return out
