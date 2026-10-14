"""박스 목록 → 고칠 수 있는 파워포인트.

글은 진짜 글상자, 네모는 도형, 그림은 원본 화소 그대로 넣는다.
화면 1920 px = 13.333 in 이므로 1 px = 6350 EMU, 글자 1 px = 0.5 pt.
"""

from __future__ import annotations

from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData, XyChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION, XL_MARKER_STYLE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

import 도표 as 도표그림
import 소재
import 스타일 as 스
from 배치 import 글머리, 글폭

EMU = 6350


def _rgb(hex색: str) -> RGBColor:
    return RGBColor.from_string(hex색.lstrip("#").upper())


# 판 맞춤: 박스는 1920×1080 기준. 다른 비율의 판(v1 템플릿 4:3 등)에서는 가로를 맞추고
# 세로는 판 높이에 맞게 늘린다(글자 크기는 가로 기준이라 줄바꿈이 그대로다. 그림은 비율을 지킨다)
_판 = {"kx": 1.0, "ky": 1.0}


def 판맞추기(prs) -> None:
    _판.update(kx=prs.slide_width / (1920 * EMU), ky=prs.slide_height / (1080 * EMU))


def _x(px: float) -> Emu:          # 가로 자리·너비
    return Emu(int(round(px * EMU * _판["kx"])))


def _y(px: float) -> Emu:          # 세로 자리·높이
    return Emu(int(round(px * EMU * _판["ky"])))


def _e(px: float) -> Emu:          # 방향 없는 길이(여백 등)
    return Emu(int(round(px * EMU * min(_판["kx"], _판["ky"]))))


def _pt(px: float) -> Pt:          # 글자 크기: 1 px = 0.5 pt (판이 작으면 같이 줄어든다)
    return Pt(px * 0.5 * _판["kx"])


def _글자체(run, 굵기: int, 꼴: dict) -> None:
    """라틴·한글 글자체를 모두 덱 글자체로. 굵기는 글꼴이 가진 것으로 맞춘다."""
    w = 스.굵기맞춤(꼴, 굵기)
    이름 = 꼴["pptx"][w]
    run.font.name = 이름
    run.font.bold = 꼴["굵게"][w]
    rPr = run._r.get_or_add_rPr()
    for 태그 in ("a:ea", "a:cs"):
        el = rPr.find(qn(태그))
        if el is None:
            el = etree.SubElement(rPr, qn(태그))
        el.set("typeface", 이름)


def _글상자(장, b: dict, 색: dict, 글자체: str) -> None:
    _, 꼴 = 스.글자체(글자체)
    tb = 장.shapes.add_textbox(_x(b["x"]), _y(b["y"]), _x(b["w"]), _y(b["h"]))
    tb.name = f"글 {b['자리']}".strip()
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = {"위": MSO_ANCHOR.TOP, "가운데": MSO_ANCHOR.MIDDLE,
                          "아래": MSO_ANCHOR.BOTTOM}[b["세로"]]
    s = b["크기"]
    들여 = 글폭(글머리, s, b["굵기"], 글자체) if b["점"] else 0
    for i, 글 in enumerate(b["글"] or [""]):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.CENTER if b["정렬"] == "가운데" else PP_ALIGN.LEFT
        p.line_spacing = _pt(s * 스.줄간격)          # 화면의 line-height 와 같은 값
        if i < len(b["글"]) - 1:
            p.space_after = _pt(0.45 * s)
        if b["점"]:
            pPr = p._p.get_or_add_pPr()
            pPr.set("marL", str(int(들여 * EMU * _판["kx"])))
            pPr.set("indent", str(-int(들여 * EMU * _판["kx"])))
            r0 = p.add_run()
            r0.text = 글머리
            r0.font.size = _pt(s)
            r0.font.color.rgb = _rgb(색["강조"])
            _글자체(r0, b["굵기"], 꼴)
        r = p.add_run()
        r.text = 글
        r.font.size = _pt(s)
        r.font.color.rgb = _rgb(색[b["색"]])
        _글자체(r, b["굵기"], 꼴)


def _네모(장, b: dict, 색: dict, 이름: str = "꾸밈"):
    모양 = {"원": MSO_SHAPE.OVAL, "화살": MSO_SHAPE.RIGHT_ARROW}.get(
        b.get("모양", ""), MSO_SHAPE.ROUNDED_RECTANGLE if b["둥글기"] else MSO_SHAPE.RECTANGLE)
    sh = 장.shapes.add_shape(모양, _x(b["x"]), _y(b["y"]), _x(b["w"]), _y(b["h"]))
    sh.name = 이름
    sh.fill.solid()
    sh.fill.fore_color.rgb = _rgb(색[b["채움"] or "강조"])
    if b.get("투명"):        # 반투명 막: <a:alpha val="불투명도×100000"/>
        clr = sh.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        etree.SubElement(clr, qn("a:alpha")).set("val", str(int((1 - b["투명"]) * 100000)))
    sh.line.fill.background()
    sh.shadow.inherit = False
    if b.get("그림자"):
        _그림자(sh)
    if 모양 == MSO_SHAPE.ROUNDED_RECTANGLE:
        sh.adjustments[0] = min(0.5, b["둥글기"] / min(b["w"], b["h"]))
    elif 모양 == MSO_SHAPE.RIGHT_ARROW:   # 화면과 같게: 머리 길이 = 높이의 절반
        sh.adjustments[0] = 0.5
        sh.adjustments[1] = 0.5
    return sh


def _그림자(sh) -> None:
    """화면의 box-shadow:0 8px 28px rgba(0,0,0,.16) 과 같은 바깥 그림자."""
    spPr = sh._element.spPr
    for el in spPr.findall(qn("a:effectLst")):
        spPr.remove(el)
    eff = etree.SubElement(spPr, qn("a:effectLst"))
    s = etree.SubElement(eff, qn("a:outerShdw"), blurRad=str(int(_e(28))), dist=str(int(_e(8))),
                         dir="5400000", algn="t", rotWithShape="0")
    c = etree.SubElement(s, qn("a:srgbClr"), val="000000")
    etree.SubElement(c, qn("a:alpha"), val="16000")


def _테두리(장, b: dict, 색: dict) -> None:
    """그림 위에 얹는 선만 있는 틀. 선 가운데가 폭/2 안쪽에 오게 해서 화면(안쪽 테두리)과 같게."""
    t = b["테두리폭"]
    모양 = MSO_SHAPE.ROUNDED_RECTANGLE if b["둥글기"] else MSO_SHAPE.RECTANGLE
    sh = 장.shapes.add_shape(모양, _x(b["x"] + t / 2), _y(b["y"] + t / 2), _x(b["w"] - t), _y(b["h"] - t))
    sh.name = "그림 테두리"
    sh.fill.background()
    sh.line.color.rgb = _rgb(색[b["테두리"]])
    sh.line.width = _e(t)
    sh.shadow.inherit = False
    if b["둥글기"]:
        sh.adjustments[0] = min(0.5, max(0, b["둥글기"] - t / 2) / min(b["w"] - t, b["h"] - t))


def _필터(pic, 필터: str) -> None:
    """원본 그림은 그대로 두고 보이는 효과만(파워포인트 '그림 서식'에서 끌 수 있다)."""
    blip = pic._element.find(".//" + qn("a:blip"))
    if 필터 == "흑백":
        etree.SubElement(blip, qn("a:grayscl"))
    elif 필터 == "옅게":
        etree.SubElement(blip, qn("a:alphaModFix"), amt="70000")


def _둥근그림(pic, 둥글기: float, w: float, h: float) -> None:
    geom = pic._element.spPr.find(qn("a:prstGeom"))
    if geom is not None and 둥글기:
        geom.set("prst", "roundRect")
        av = geom.find(qn("a:avLst"))
        if av is None:
            av = etree.SubElement(geom, qn("a:avLst"))
        gd = etree.SubElement(av, qn("a:gd"))
        gd.set("name", "adj")
        gd.set("fmla", f"val {int(min(50000, 둥글기 / min(w, h) * 100000))}")


def _그림(장, b: dict, 색: dict, 프로젝트: Path, 알림: list) -> None:
    경로 = 프로젝트 / b["그림"]
    if not 경로.exists():
        알림.append(f"그림을 못 찾았어요: {b['그림']}")
        return
    크기 = 소재.크기안전(경로)
    if 크기 is None:
        알림.append(f"그림으로 읽지 못해 빼고 구웠어요: {b['그림']}")
        return
    if b["채움"]:
        _네모(장, {**b, "종류": "네모", "모양": ""}, 색, "그림 바탕")
    iw, ih = 크기
    try:
        넣을것, 말 = 소재.파워포인트용(경로)
    except Exception:  # noqa: BLE001
        알림.append(f"그림을 PPTX에 넣지 못해 빼고 구웠어요: {b['그림']}")
        return
    if 말:
        알림.append(말)
    # 판 위의 실제 상자(EMU)에서 비율을 계산해야 그림이 찌그러지지 않는다
    X, Y, W, H = _x(b["x"]), _y(b["y"]), _x(b["w"]), _y(b["h"])
    if b["맞춤"] == "담기":                     # 전체가 보이게 줄여서 가운데
        k = min(W / iw, H / ih)
        pw, ph = int(iw * k), int(ih * k)
        pic = 장.shapes.add_picture(넣을것, Emu(X + (W - pw) // 2), Emu(Y + (H - ph) // 2),
                                   Emu(pw), Emu(ph))
    else:                                       # 꽉 채우고 넘치는 쪽을 잘라 냄(원본은 그대로)
        pic = 장.shapes.add_picture(넣을것, X, Y, W, H)
        상자비, 그림비 = W / H, iw / ih
        if 그림비 > 상자비:
            c = (1 - 상자비 / 그림비) / 2
            pic.crop_left = pic.crop_right = c
        else:
            c = (1 - 그림비 / 상자비) / 2
            pic.crop_top = pic.crop_bottom = c
        _둥근그림(pic, b["둥글기"], b["w"], b["h"])
        if b.get("그림자") and not b["채움"]:
            _그림자(pic)
    pic.name = f"그림 {Path(b['그림']).name}"
    if b.get("필터"):
        _필터(pic, b["필터"])
    if b.get("테두리"):
        _테두리(장, b, 색)


def _도표(장, b: dict, 색: dict, 글자체: str) -> None:
    """파워포인트 기본 차트: 데이터가 차트 안 엑셀에 들어가서 파워포인트에서 값을 고칠 수 있다."""
    d = b["도표"]
    if not d or not d.get("계열"):
        return
    _, 꼴 = 스.글자체(글자체)
    계열색 = 색.get("계열") or [색["강조"]]
    if d["모양"] == "점":
        cd = XyChartData()
        for s in d["계열"]:
            ser = cd.add_series(s["이름"])
            for x, y in zip(d["가로"], s["값"]):
                if y is not None:
                    ser.add_data_point(float(x), y)
        종류 = XL_CHART_TYPE.XY_SCATTER
    else:
        cd = CategoryChartData()
        cd.categories = [str(c) for c in d["가로"]]
        for s in d["계열"]:
            cd.add_series(s["이름"], s["값"])
        종류 = XL_CHART_TYPE.COLUMN_CLUSTERED if d["모양"] == "막대" else XL_CHART_TYPE.LINE_MARKERS
    gf = 장.shapes.add_chart(종류, _x(b["x"]), _y(b["y"]), _x(b["w"]), _y(b["h"]), cd)
    gf.name = "도표"
    ch = gf.chart
    ch.font.size = _pt(26)
    ch.font.name = 꼴["pptx"][400]
    ch.font.color.rgb = _rgb(색["흐린글"])
    ch.has_legend = len(d["계열"]) > 1
    if ch.has_legend:
        ch.legend.position = XL_LEGEND_POSITION.TOP
        ch.legend.include_in_layout = False
        ch.legend.font.size = _pt(28)
        ch.legend.font.color.rgb = _rgb(색["글"])
    va = ch.value_axis
    va.has_major_gridlines = True
    va.major_gridlines.format.line.color.rgb = _rgb(색["선"])
    va.format.line.fill.background()
    # 화면(SVG)과 같은 눈금: 범위·간격을 똑같이 정한다
    눈금 = 도표그림.눈금(*도표그림.범위(d))
    va.minimum_scale, va.maximum_scale = 눈금[0], 눈금[-1]
    va.major_unit = 눈금[1] - 눈금[0]
    ca = ch.category_axis
    ca.format.line.color.rgb = _rgb(색["선"])
    ca.has_title = True
    ca.axis_title.text_frame.text = str(d["가로이름"])
    ca.axis_title.text_frame.paragraphs[0].runs[0].font.size = _pt(26)
    for j, plot_s in enumerate(ch.plots[0].series):
        c = _rgb(계열색[j % len(계열색)])
        if d["모양"] == "막대":
            plot_s.format.fill.solid()
            plot_s.format.fill.fore_color.rgb = c
        else:
            plot_s.format.line.color.rgb = c
            plot_s.format.line.width = Pt(2.5)
            plot_s.marker.style = XL_MARKER_STYLE.CIRCLE
            plot_s.marker.size = 8
            plot_s.marker.format.fill.solid()
            plot_s.marker.format.fill.fore_color.rgb = c
            if d["모양"] == "점":
                plot_s.format.line.fill.background()
    if d["모양"] == "막대":
        ch.plots[0].gap_width = 60
        ch.plots[0].overlap = -8
        n = len(d["가로"]) * len(d["계열"])
        if n <= 16:   # 막대가 적으면 값 표시
            ch.plots[0].has_data_labels = True
            dl = ch.plots[0].data_labels
            dl.font.size = _pt(24)
            dl.font.color.rgb = _rgb(색["글"])


def _아래선(cell, 색값: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    ln = etree.SubElement(tcPr, qn("a:lnB"), w=str(int(_e(2))), cap="flat", cmpd="sng")
    sf = etree.SubElement(ln, qn("a:solidFill"))
    etree.SubElement(sf, qn("a:srgbClr"), val=색값.lstrip("#").upper())
    tcPr.insert(0, ln)        # 선은 칸 채움보다 앞에 와야 한다(스키마 차례)


def _표(장, b: dict, 색: dict, 글자체: str) -> None:
    줄들 = b["표"]
    if not 줄들 or not 줄들[0]:
        return
    _, 꼴 = 스.글자체(글자체)
    n줄, n칸 = len(줄들), max(len(r) for r in 줄들)
    gf = 장.shapes.add_table(n줄, n칸, _x(b["x"]), _y(b["y"]), _x(b["w"]), _y(b["h"]))
    gf.name = "표"
    t = gf.table
    # 기본 표 스타일(흰 격자선·띠)을 빼고 화면처럼 칸 색만 쓴다: '스타일 없음, 격자 없음'
    tblPr = gf._element.graphic.graphicData.tbl.tblPr
    tblPr.set("firstRow", "0")
    tblPr.set("bandRow", "0")
    sid = tblPr.find(qn("a:tableStyleId"))
    if sid is None:
        sid = etree.SubElement(tblPr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"
    for j, 너비 in enumerate(b.get("칸너비들") or []):   # 칸 너비는 화면과 같은 값
        if j < n칸:
            t.columns[j].width = _x(너비)
    정렬들 = b.get("정렬들") or []
    for i, 줄 in enumerate(줄들):
        t.rows[i].height = _y(b["h"] / n줄)
        for j in range(n칸):
            cell = t.cell(i, j)
            cell.fill.solid()
            cell.fill.fore_color.rgb = _rgb(색["강조"] if i == 0 else (색["면"] if i % 2 == 0 else 색["바탕"]))
            cell.margin_left = cell.margin_right = _e(18)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf = cell.text_frame
            tf.word_wrap = False                         # 화면처럼 한 줄(넘치면 감독관이 잡는다)
            tf.text = str(줄[j]) if j < len(줄) else ""
            if j < len(정렬들) and 정렬들[j] == "오른":
                tf.paragraphs[0].alignment = PP_ALIGN.RIGHT
            if i == n줄 - 1:                             # 화면의 표 아래 2px 선
                _아래선(cell, 색["선"])
            for r in tf.paragraphs[0].runs:
                r.font.size = _pt(b["크기"])
                r.font.color.rgb = _rgb(색["바탕"] if i == 0 else 색["글"])
                _글자체(r, 700 if i == 0 else 400, 꼴)


def 장그리기(장, 장자료: dict, 프로젝트: Path, 글자체: str | None, 알림: list,
          바탕칠: bool = True) -> None:
    """박스들을 장 하나에 그린다. 바탕칠=False 면 템플릿(마스터)의 배경을 그대로 둔다."""
    색 = 장자료["색"]
    if 바탕칠:
        장.background.fill.solid()
        장.background.fill.fore_color.rgb = _rgb(색["바탕"])
    for b in 장자료["박스"]:
        if b["종류"] == "글":
            _글상자(장, b, 색, 글자체)
        elif b["종류"] == "그림":
            _그림(장, b, 색, 프로젝트, 알림)
        elif b["종류"] == "도표":
            _도표(장, b, 색, 글자체)
        elif b["종류"] == "표":
            _표(장, b, 색, 글자체)
        else:
            _네모(장, b, 색)
    if 장자료.get("노트"):
        장.notes_slide.notes_text_frame.text = 장자료["노트"]


def 만들기(장들: list[dict], 프로젝트: Path, 저장경로: Path, 글자체: str | None = None) -> list[str]:
    """장들: [{"박스": [...], "색": {...}, "노트": "..."}] → 새 .pptx. 알릴 말 목록을 돌려준다."""
    prs = Presentation()
    prs.slide_width, prs.slide_height = Emu(1920 * EMU), Emu(1080 * EMU)
    판맞추기(prs)
    빈판 = prs.slide_layouts[6]
    알림: list[str] = []
    for 장자료 in 장들:
        장그리기(prs.slides.add_slide(빈판), 장자료, 프로젝트, 글자체, 알림)
    저장경로.parent.mkdir(parents=True, exist_ok=True)
    prs.save(저장경로)
    return 알림
