"""덱편집기 1단계 시험: 박스 모델, 화면·PPTX 일치, 굽기, 서버."""

from __future__ import annotations

import json
import shutil
import sys
import threading
import urllib.request
from http.server import ThreadingHTTPServer
from pathlib import Path

import pytest
from pptx import Presentation

앱 = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(앱))

import 굽기  # noqa: E402
import 배치  # noqa: E402
import 엔진  # noqa: E402
import 그리기_pptx  # noqa: E402

본보기 = 앱 / "프로젝트" / "본보기_양극재수세"


@pytest.fixture
def 프로젝트(tmp_path, monkeypatch):
    """본보기를 임시 폴더에 복사해서 쓴다(원본은 건드리지 않는다)."""
    밑 = tmp_path / "프로젝트"
    shutil.copytree(본보기, 밑 / "시험덱", ignore=shutil.ignore_patterns("내보내기", ".이전", ".미리보기"))
    monkeypatch.setattr(엔진, "프로젝트들", 밑)
    return "시험덱"


def test_글맞추기_줄이고_넘침표시():
    b = 배치.글맞추기(배치.박스("글", 0, 0, 600, 120, 글=["짧은 글"], 크기=60))
    assert b.크기 == 60 and not b.넘침
    긴글 = ["아주 긴 문장이 여러 번 되풀이됩니다 " * 6]
    b = 배치.글맞추기(배치.박스("글", 0, 0, 600, 120, 글=긴글, 크기=60))
    assert b.크기 == 배치.스.글자["가장작은"] and b.넘침


def test_줄나누기는_띄어쓰기_단위():
    줄들 = 배치.줄나누기("수세 뒤 1차 입자는 더 둥글고 고르게 바뀌었다", 400, 40)
    assert len(줄들) >= 2
    assert all(not 줄.startswith(" ") for 줄 in 줄들)
    assert "".join(줄들).replace(" ", "") == "수세뒤1차입자는더둥글고고르게바뀌었다"


def test_장종류마다_박스(프로젝트):
    설계 = 엔진.읽기(프로젝트)
    덱 = 엔진.덱배치(설계, 엔진.폴더(프로젝트))
    assert [장["종류"] for 장 in 덱["장들"]] == [장["종류"] for 장 in 설계["장"]]
    for 장 in 덱["장들"]:
        assert 장["박스"] and 장["넘침"] == 0
        for b in 장["박스"]:
            assert 0 <= b["x"] and b["x"] + b["w"] <= 1920 + 1
            assert 0 <= b["y"] and b["y"] + b["h"] <= 1080 + 1


def test_pptx_글자가_살아있고_자리가_같다(프로젝트, tmp_path):
    덱 = 엔진.덱배치(엔진.읽기(프로젝트), 엔진.폴더(프로젝트))
    f = tmp_path / "덱.pptx"
    알림 = 그리기_pptx.만들기([{**장, "색": 덱["색"]} for 장 in 덱["장들"]], 엔진.폴더(프로젝트), f,
                         덱["글자체"])
    assert 알림 == []
    prs = Presentation(f)
    assert len(prs.slides) == len(덱["장들"])
    assert prs.slide_width == 1920 * 6350
    for 장, 장자료 in zip(prs.slides, 덱["장들"]):
        글상자 = [s for s in 장.shapes if s.has_text_frame and s.name.startswith("글")]
        글박스 = [b for b in 장자료["박스"] if b["종류"] == "글"]
        assert len(글상자) == len(글박스)
        for s, b in zip(글상자, 글박스):
            assert s.left == round(b["x"] * 6350) and s.top == round(b["y"] * 6350)
            assert s.text_frame.text.replace("• ", "").replace("\n", "") == "".join(b["글"])
            run = s.text_frame.paragraphs[0].runs[-1]
            assert run.font.name == "맑은 고딕"
            assert abs(run.font.size.pt - b["크기"] / 2) < 0.01
    그림수 = sum(1 for 장 in prs.slides for s in 장.shapes if s.shape_type == 13)
    assert 그림수 == sum(b["종류"] == "그림" for 장 in 덱["장들"] for b in 장["박스"])
    assert "수세 공정" in prs.slides[0].notes_slide.notes_text_frame.text


def test_tiff원본_화소가_그대로_들어간다(프로젝트, tmp_path):
    import io
    import numpy as np
    from PIL import Image
    설계 = 엔진.읽기(프로젝트)
    설계["장"] = [{"종류": "글그림", "제목": "그림", "그림": "소재/SEM_1차입자_x31750.tif"}]
    덱 = 엔진.덱배치(설계, 엔진.폴더(프로젝트))
    f = tmp_path / "덱.pptx"
    그리기_pptx.만들기([{**장, "색": 덱["색"]} for 장 in 덱["장들"]], 엔진.폴더(프로젝트), f)
    pic = next(s for s in Presentation(f).slides[0].shapes if s.shape_type == 13)
    넣은것 = np.asarray(Image.open(io.BytesIO(pic.image.blob)))
    원본 = np.asarray(Image.open(엔진.폴더(프로젝트) / "소재" / "SEM_1차입자_x31750.tif"))
    assert 넣은것.shape == 원본.shape and np.array_equal(넣은것, 원본)


@pytest.mark.skipif(not any(Path(p).exists() for p in 굽기._후보 if p), reason="브라우저 없음")
def test_굽기_png_pdf(프로젝트):
    from PIL import Image
    r = 엔진.굽기실행(프로젝트, ["pptx", "png", "pdf"])
    이름들 = [f["이름"] for f in r["파일"]]
    assert [n.rsplit(".", 1)[1] for n in 이름들] == ["pptx", "pdf", "zip"]
    pngs = sorted((엔진.폴더(프로젝트) / r["폴더"] / "그림").glob("*.png"))
    assert len(pngs) == len(엔진.읽기(프로젝트)["장"])
    for p in pngs:
        im = Image.open(p)
        assert im.size == (1920, 1080)
        # 아래쪽이 잘려 검게 비지 않았나(예전 크롬 스크린샷 문제)
        assert im.convert("L").crop((0, 1000, 1920, 1080)).getextrema()[1] > 200


def test_서버_길(프로젝트):
    import 서버
    srv = ThreadingHTTPServer(("127.0.0.1", 0), 서버.처리기)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    밑 = f"http://127.0.0.1:{srv.server_address[1]}"

    def 묻기(길, 몸):
        req = urllib.request.Request(밑 + "/api/" + urllib.request.quote(길),
                                     data=json.dumps(몸).encode(), method="POST",
                                     headers={"Content-Type": "application/json"})
        with urllib.request.urlopen(req) as r:
            return json.loads(r.read())

    try:
        assert 묻기("목록", {})["프로젝트"][0]["이름"] == 프로젝트
        자료 = 묻기("열기", {"이름": 프로젝트})
        assert len(자료["장들"]) == len(자료["설계"]["장"]) and "<section" in 자료["장들"][0]["html"]
        자료["설계"]["장"][1]["큰말"] = "바뀐 한마디"
        자료 = 묻기("저장", {"이름": 프로젝트, "설계": 자료["설계"]})
        assert "바뀐 한마디" in 자료["장들"][1]["html"]
        assert list((엔진.폴더(프로젝트) / ".이전").glob("덱설계_*.json"))   # 이전 판을 남겼다
        with urllib.request.urlopen(밑 + "/" + urllib.request.quote("장.css")) as r:
            assert b".\xec\x9e\xa5" in r.read()     # ".장" CSS
    finally:
        srv.shutdown()


# ---------------------------------------------------------------- 3단계: 재료·도표·표
import 재료받기  # noqa: E402

받은자료 = 본보기 / "받은자료"


def test_pdf_회의록에서_글덩어리와_그림(tmp_path):
    m = 재료받기.넣기(tmp_path, 받은자료 / "A사_미팅회의록_20261008.pdf")
    글들 = [g["글"] for g in m["글"]]
    assert m["쪽수"] == 2 and len(m["쪽"]) == 2
    assert "고객사 미팅 회의록" in 글들
    # 줄바꿈으로 잘린 문단이 하나로 이어졌다
    assert any(g.startswith("A사는 수세 공정") and g.endswith("싶어 했다.") for g in 글들)
    assert any(g.startswith("A사: 500사이클") for g in 글들)
    확장 = sorted(Path(p).suffix for p in m["그림"])
    assert 확장 == [".jpg", ".png"]                     # JPEG는 JPEG 그대로(다시 압축하지 않음)


def test_엑셀_표(tmp_path):
    m = 재료받기.넣기(tmp_path, 받은자료 / "수세실험_데이터.xlsx")
    표 = {t["이름"]: t for t in m["표"]}
    assert 표["수명"]["머리"] == ["사이클", "수세 전 (%)", "수세 후 (%)"] and 표["수명"]["줄수"] == 6


def test_pptx_재료는_원본그림_그대로(tmp_path, 프로젝트):
    덱 = 엔진.덱배치(엔진.읽기(프로젝트), 엔진.폴더(프로젝트))
    f = tmp_path / "v1.pptx"
    그리기_pptx.만들기([{**장, "색": 덱["색"]} for 장 in 덱["장들"]], 엔진.폴더(프로젝트), f)
    m = 재료받기.넣기(tmp_path / "p", f)
    assert m["쪽수"] == len(덱["장들"]) and any("수세 공정" in g["글"] for g in m["글"] if g.get("자리") == "노트")
    blobs = {s.image.blob for 장 in Presentation(f).slides for s in 장.shapes if s.shape_type == 13}
    assert {(tmp_path / "p" / g).read_bytes() for g in m["그림"]} == blobs


def test_도표장_기본차트와_표(프로젝트, tmp_path):
    p = 엔진.폴더(프로젝트)
    엑셀 = 재료받기.넣기(p, 받은자료 / "수세실험_데이터.xlsx")
    표 = {t["이름"]: t["경로"] for t in 엑셀["표"]}
    설계 = 엔진.읽기(프로젝트)
    설계["장"] = [{"종류": "도표", "제목": "수명", "표": 표["수명"], "모양": "선"},
                 {"종류": "도표", "제목": "입도", "표": 표["입도분포"], "모양": "막대"},
                 {"종류": "표", "제목": "입도 표", "표": 표["입도분포"]}]
    덱 = 엔진.덱배치(설계, p)
    선 = next(b for b in 덱["장들"][0]["박스"] if b["종류"] == "도표")["도표"]
    assert 선["가로"] == ["0", "100", "200", "300", "400", "500"] and len(선["계열"]) == 2
    f = tmp_path / "도표.pptx"
    그리기_pptx.만들기([{**장, "색": 덱["색"]} for 장 in 덱["장들"]], p, f, 덱["글자체"])
    prs = Presentation(f)
    ch = next(s for s in prs.slides[0].shapes if s.has_chart).chart
    assert [s.name for s in ch.plots[0].series] == ["수세 전 (%)", "수세 후 (%)"]
    assert list(ch.plots[0].series[1].values) == [100, 98, 96, 95, 93, 92]   # 차트 안 데이터 = 엑셀 값
    assert ch.value_axis.minimum_scale == 70                                 # 화면과 같은 눈금
    막대 = next(s for s in prs.slides[1].shapes if s.has_chart).chart
    assert 막대.value_axis.minimum_scale == 0                                # 막대는 0에서
    t = next(s for s in prs.slides[2].shapes if s.has_table).table
    assert t.cell(0, 1).text == "D10 (µm)" and t.cell(3, 0).text == "수세 후 (1:2)"
    html = 엔진.화면자료(프로젝트, 설계)["장들"][0]["html"]
    assert "<svg" in html and "<polyline" in html


@pytest.mark.skipif(not 재료받기.ffmpeg(), reason="ffmpeg 없음")
def test_영상_장면뽑기(tmp_path):
    import subprocess
    v = tmp_path / "v.mp4"
    subprocess.run([재료받기.ffmpeg(), "-loglevel", "error", "-f", "lavfi", "-i",
                    "color=c=red:s=320x240:d=1", "-f", "lavfi", "-i", "color=c=blue:s=320x240:d=1",
                    "-filter_complex", "[0][1]concat=n=2:v=1", "-pix_fmt", "yuv420p", str(v)], check=True)
    m = 재료받기.넣기(tmp_path / "p", v)
    assert len(m["그림"]) == 2


# ---------------------------------------------------------------- 2단계: 장 종류 16 + 스타일 10
모음 = 앱 / "프로젝트" / "본보기_장종류모음"


@pytest.fixture
def 모음덱(tmp_path, monkeypatch):
    밑 = tmp_path / "프로젝트"
    shutil.copytree(모음, 밑 / "모음", ignore=shutil.ignore_patterns("내보내기", ".이전", ".미리보기"))
    monkeypatch.setattr(엔진, "프로젝트들", 밑)
    return "모음"


def test_장종류_16가지_스타일_10개_모두_넘침없음(모음덱):
    import 스타일
    설계 = 엔진.읽기(모음덱)
    종류들 = {장["종류"] for 장 in 설계["장"]}
    assert 종류들 == set(배치.장종류) - {"원본장"} and len(종류들) == 16
    assert len(스타일.스타일들) == 10
    for 이름, st in 스타일.스타일들.items():
        덱 = 엔진.덱배치({**설계, "스타일": 이름}, 엔진.폴더(모음덱))
        for 장 in 덱["장들"]:
            assert 장["넘침"] == 0, (이름, 장["종류"])
            for b in 장["박스"]:
                assert 0 <= b["x"] and b["x"] + b["w"] <= 1921 and 0 <= b["y"] and b["y"] + b["h"] <= 1081
                for 역할 in (b["색"], b["채움"]):
                    assert not 역할 or 역할 in 덱["색"], (이름, 역할)


def test_스타일_장식이_배치를_바꾼다():
    장 = {"종류": "표지", "큰말": "제목", "그림": "a.png"}
    기본 = 배치.장배치(장)
    띠 = 배치.장배치(장, 장식값={"표지": "띠"})
    배경 = 배치.장배치(장, 장식값={"표지": "그림배경"})
    assert 기본[-1].종류 == "그림" and 기본[-1].x == 1110
    assert 띠[0].채움 == "강조" and 띠[0].w == 1100
    assert 배경[0].종류 == "그림" and 배경[0].w == 1920 and 배경[1].투명 > 0


def test_반투명_화살표가_pptx_도형으로(모음덱, tmp_path):
    설계 = 엔진.읽기(모음덱)
    for 장 in 설계["장"]:
        if 장["종류"] == "사진가득":
            장["맞춤"] = "채우기"             # 꽉 채우기일 때 반투명 띠가 깔린다
    덱 = 엔진.덱배치(설계, 엔진.폴더(모음덱))
    f = tmp_path / "모음.pptx"
    그리기_pptx.만들기([{**장, "색": 덱["색"]} for 장 in 덱["장들"]], 엔진.폴더(모음덱), f, 덱["글자체"])
    prs = Presentation(str(f))
    assert len(prs.slides) == len(덱["장들"])
    xml = "".join(s._element.xml for s in prs.slides)
    assert 'prst="rightArrow"' in xml and 'prst="ellipse"' in xml and "<a:alpha " in xml


# ---------------------------------------------------------------- 해부
def 참고덱만들기(f: Path) -> Path:
    """어두운 바탕 + 금색 표지 띠 + 제목 왼쪽 세로 띠인 참고 PPTX."""
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Emu as E
    prs = Presentation()
    prs.slide_width, prs.slide_height = E(12192000), E(6858000)
    W, H = prs.slide_width, prs.slide_height
    for i in range(3):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        s.background.fill.solid()
        s.background.fill.fore_color.rgb = RGBColor(0x10, 0x18, 0x26)
        if i == 0:
            띠 = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, int(W * 0.45), H)
        else:
            띠 = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, int(W * 0.05), int(H * 0.08), int(W * 0.008), int(H * 0.08))
        띠.fill.solid()
        띠.fill.fore_color.rgb = RGBColor(0xE0, 0xA5, 0x26)
        t = s.shapes.add_textbox(int(W * 0.07), int(H * 0.07), int(W * 0.8), int(H * 0.1))
        r = t.text_frame.paragraphs[0].add_run()
        r.text = f"참고 장 {i + 1}"
        r.font.name = "Pretendard"
    prs.save(f)
    return f


def test_해부_색과_장식을_읽는다(tmp_path):
    import 해부
    s = 해부.해부(참고덱만들기(tmp_path / "좋은발표_v3.pptx"))
    assert s["이름"] == "좋은발표" and s["글자체"] == "Pretendard"
    assert s["색"]["바탕"] == "#101826" and s["색"]["강조"] == "#E0A526"
    assert 배치.스.색갖추기(s)["반전글"] in ("#111111", "#101826")     # 금색 위에는 어두운 글
    assert s["장식"]["제목"] == "왼띠" and s["장식"]["표지"] == "띠"
    acme = 해부.해부(본보기 / "받은자료" / "ACME_중간보고_v1.pptx")
    assert acme["색"]["강조"] == "#C8102E" and acme["장식"]["제목"] == "위띠"


def test_해부하기_덱에_입히고_남긴다(프로젝트, tmp_path, monkeypatch):
    import base64
    import 해부
    import 스타일
    monkeypatch.setattr(스타일, "내스타일폴더", tmp_path / "내스타일")
    monkeypatch.setattr(해부, "내스타일폴더", tmp_path / "내스타일")
    f = 참고덱만들기(tmp_path / "좋은발표.pptx")
    원래 = f.read_bytes()
    s = 엔진.해부하기(프로젝트, f.name, base64.b64encode(원래).decode())
    s2 = 엔진.해부하기(프로젝트, f.name, base64.b64encode(원래).decode())
    assert s["이름"] == "좋은발표" and s2["이름"] == "좋은발표 (2)"        # 덮어쓰지 않는다
    assert f.read_bytes() == 원래
    설계 = 엔진.읽기(프로젝트)
    assert 설계["스타일"]["색"]["강조"] == "#E0A526"
    덱 = 엔진.덱배치(설계, 엔진.폴더(프로젝트))
    assert 덱["색"]["바탕"] == "#101826" and all(장["넘침"] == 0 for 장 in 덱["장들"])
    assert "좋은발표" in 엔진.화면자료(프로젝트)["내스타일들"]


# ---------------------------------------------------------------- 4단계: 감독관
def test_감독관_시험은_0틀림():
    import 감독관
    assert 감독관.시험() == 0


def test_멈춤이면_굽지_않고_강행하면_굽는다(프로젝트):
    설계 = 엔진.읽기(프로젝트)
    설계["장"].append({"종류": "한마디", "큰말": "한 문장으로 말할 결론을 넣으세요"})
    엔진.쓰기(프로젝트, 설계)
    r = 엔진.굽기실행(프로젝트, ["pptx"])
    assert r["멈춤"] and r["멈춤"][0]["규칙"] == "견본글" and not r["파일"]
    assert not (엔진.폴더(프로젝트) / "내보내기").exists()
    r = 엔진.굽기실행(프로젝트, ["pptx"], 강행=True)
    assert r["파일"] and r["점검"] == []


def test_없는_그림도_덱은_열리고_고침은_지우지_않는다(프로젝트):
    설계 = 엔진.읽기(프로젝트)
    설계["장"].append({"종류": "글그림", "제목": "단면", "글": ["틈이 보인다"], "그림": "소재/없는.tif"})
    엔진.쓰기(프로젝트, 설계)
    자료 = 엔진.화면자료(프로젝트)
    n = len(설계["장"])
    찾음 = [f for f in 자료["감독"] if f["규칙"] == "그림없음"]
    assert 찾음 and 찾음[0]["장"] == n and 찾음[0]["고침"] == "글만으로"
    자료 = 엔진.감독고치기(프로젝트, "그림없음", n, "글만으로")
    설계 = 엔진.읽기(프로젝트)
    assert 설계["장"][-1]["종류"] == "글만" and 설계["휴지통"][-1]["그림"] == "소재/없는.tif"
    assert not [f for f in 자료["감독"] if f["규칙"] == "그림없음"]
    자료 = 엔진.감독무시(프로젝트, "그림되풀이", 1)
    assert 엔진.읽기(프로젝트)["감독관무시"][-1]["규칙"] == "그림되풀이"


# ---------------------------------------------------------------- 5단계: 꾸미기
def test_꾸밈_기본은_없음_그리고_고르면_화면과_PPTX에_같이(프로젝트, tmp_path):
    import io
    import numpy as np
    from PIL import Image
    설계 = 엔진.읽기(프로젝트)
    설계["장"] = [{"종류": "글그림", "제목": "그림", "글": ["둥글다"], "그림": "소재/SEM_1차입자_x31750.tif"}]
    기본 = 엔진.덱배치(설계, 엔진.폴더(프로젝트))
    assert not any(b.get("테두리") or b.get("그림자") or b.get("필터") for b in 기본["장들"][0]["박스"])
    설계["꾸밈"] = {"배경": "아래띠", "사진": "흑백", "테두리": "강조", "그림자": "부드럽게", "모서리": "각지게"}
    덱 = 엔진.덱배치(설계, 엔진.폴더(프로젝트))
    박스들 = 덱["장들"][0]["박스"]
    assert len(박스들) == len(기본["장들"][0]["박스"]) + 1          # 아래띠 하나
    그림 = next(b for b in 박스들 if b["종류"] == "그림")
    assert 그림["필터"] == "흑백" and 그림["테두리"] == "강조" and 그림["그림자"] and 그림["둥글기"] == 0
    html = 엔진.화면.장html(박스들, 덱["색"])
    assert "grayscale(1)" in html and "box-shadow" in html
    f = tmp_path / "꾸밈.pptx"
    그리기_pptx.만들기([{**장, "색": 덱["색"]} for 장 in 덱["장들"]], 엔진.폴더(프로젝트), f)
    s = Presentation(f).slides[0]
    xml = s._element.xml
    assert "<a:grayscl/>" in xml and "<a:outerShdw" in xml
    assert any(sh.name == "그림 테두리" for sh in s.shapes)
    pic = next(sh for sh in s.shapes if sh.shape_type == 13)          # 효과만, 원본 화소는 그대로
    원본 = np.asarray(Image.open(엔진.폴더(프로젝트) / "소재" / "SEM_1차입자_x31750.tif"))
    assert np.array_equal(np.asarray(Image.open(io.BytesIO(pic.image.blob))), 원본)


def test_꽉찬_장에는_배경무늬를_깔지_않는다():
    import 꾸밈
    박스들 = [b.사전() for b in 배치.장배치({"종류": "끝", "큰말": "감사합니다"})]
    assert 꾸밈.입히기(박스들, {"배경": "점무늬"}) == 박스들


# ---------------------------------------------------------------- 6단계: 내보내기 전부 + 한 명령
@pytest.mark.skipif(not any(Path(p).exists() for p in 굽기._후보 if p), reason="브라우저 없음")
def test_덱만들기_한_명령으로_전부(프로젝트, capsys):
    import hashlib
    import zipfile
    import 덱만들기
    import 재료받기
    폴더 = 엔진.폴더(프로젝트)
    설계 = 엔진.읽기(프로젝트)
    설계["장"][0]["시간"] = 3
    엔진.쓰기(프로젝트, 설계)
    assert 덱만들기.main([str(폴더 / "덱설계.json")]) == 0
    out = next((폴더 / "내보내기").iterdir())
    n = len(설계["장"])
    이름 = 설계["이름"]
    assert len(Presentation(out / f"{이름}.pptx").slides) == n
    assert len(zipfile.ZipFile(out / f"{이름}_그림묶음.zip").namelist()) == n
    assert "## 1장" in (out / "대본.md").read_text("utf-8")
    from openpyxl import load_workbook
    assert len(load_workbook(out / "데이터.xlsx").sheetnames) >= 2
    for f in (out / "소재").iterdir():                               # 원본 그대로
        assert hashlib.md5(f.read_bytes()).hexdigest() == hashlib.md5((폴더 / "소재" / f.name).read_bytes()).hexdigest()
    if 재료받기.ffmpeg():
        import subprocess
        길이 = float(subprocess.run(["ffprobe", "-v", "error", "-show_entries", "format=duration", "-of", "csv=p=0",
                                   str(out / f"{이름}.mp4")], capture_output=True, text=True).stdout)
        assert abs(길이 - (3 + 5 * (n - 1) + 0.5)) < 0.2
    assert "멈춤" not in capsys.readouterr().out


def test_덱만들기_멈춤이면_끝코드_2(프로젝트):
    import 덱만들기
    설계 = 엔진.읽기(프로젝트)
    설계["장"].append({"종류": "한마디", "큰말": "결론을 넣으세요"})
    엔진.쓰기(프로젝트, 설계)
    assert 덱만들기.main([프로젝트, "--형식", "pptx"]) == 2
    assert not (엔진.폴더(프로젝트) / "내보내기").exists()


# ---------------------------------------------------------------- 7단계: Claude용 쓰는 법·스키마·첫 실전 덱
def test_스키마와_쓰는법이_앱의_장종류와_같다():
    import 설계틀
    assert json.loads(설계틀.파일.read_text("utf-8")) == 설계틀.스키마(), "python 설계틀.py 를 다시 돌려 주세요"
    assert 설계틀.쓰는법.read_text("utf-8") == 설계틀.쓰는법글(), "python 설계틀.py 를 다시 돌려 주세요"


@pytest.mark.parametrize("덱", ["A사_2차미팅_수세조건", "본보기_양극재수세", "본보기_장종류모음"])
def test_본보기_덱은_스키마를_지키고_감독관이_멈추지_않는다(덱):
    jsonschema = pytest.importorskip("jsonschema")
    import 감독관
    import 설계틀
    p = 앱 / "프로젝트" / 덱
    설계 = json.loads((p / "덱설계.json").read_text("utf-8"))
    jsonschema.validate(설계, 설계틀.스키마())
    찾음 = 감독관.살피기(설계, 엔진.덱배치(설계, p), p)
    assert not [f for f in 찾음 if f["수준"] in ("멈춤", "주의")], 찾음


def test_재료넣기_명령은_새덱을_만들고_요약을_쓴다(tmp_path, monkeypatch):
    import 재료넣기
    monkeypatch.setattr(엔진, "프로젝트들", tmp_path / "프로젝트")
    assert 재료넣기.main(["새 실험", str(받은자료 / "수세실험_데이터.xlsx"),
                        str(받은자료 / "실험메모_수세_2차.md")]) == 0
    요약 = (tmp_path / "프로젝트" / "새_실험" / "재료요약.md").read_text("utf-8")
    assert "재료/001_수세실험_데이터/표/수명.json" in 요약 and "| 500 | 77 | 92 |" in 요약
    assert "300사이클부터" in 요약


# ---------------------------------------------------------------- 마지막 점검에서 고친 것
def test_표_숫자는_반올림하지_않고_열마다_자리를_맞춘다():
    b = [x for x in 배치.장배치({"종류": "표", "제목": "t", "_표": {"머리": ["시료", "값", "비"], "줄들": [
        ["가", 1234.5, 6.1], ["나", 89.95, 10.0], ["다", 10.0, None]]}}) if x.종류 == "표"][0]
    assert [r[1] for r in b.표[1:]] == ["1,234.50", "89.95", "10.00"]
    assert [r[2] for r in b.표[1:]] == ["6.1", "10.0", ""]
    assert b.정렬들 == ["왼", "오른", "오른"] and abs(sum(b.칸너비들) - b.w) < 1


def test_SEM은_기본으로_잘리지_않는다():
    for 장 in ({"종류": "표지", "큰말": "t", "그림": "a.tif"}, {"종류": "사진가득", "제목": "t", "그림": "a.tif"}):
        그림 = [b for b in 배치.장배치(장) if b.종류 == "그림"]
        assert 그림 and all(b.맞춤 == "담기" for b in 그림)


def test_깨진_그림과_이상한_도표에도_덱이_열린다(프로젝트):
    p = 엔진.폴더(프로젝트)
    (p / "소재" / "빈.png").write_bytes(b"")
    설계 = 엔진.읽기(프로젝트)
    설계["장"].append({"종류": "사진가득", "제목": "깨짐", "그림": "소재/빈.png"})
    for 장 in 설계["장"]:
        if 장["종류"] == "도표":
            장["모양"] = "점"                    # 가로가 글자인 표로 점 도표
    엔진.쓰기(프로젝트, 설계)
    자료 = 엔진.화면자료(프로젝트)
    assert any(f["규칙"] == "그림읽기" for f in 자료["감독"])
    r = 엔진.굽기실행(프로젝트, ["pptx"], 강행=True)
    assert r["파일"] and any("빈.png" in m for m in r["알림"])


def test_서버는_다른_사이트의_요청을_막는다():
    import http.client
    import 서버
    srv = ThreadingHTTPServer(("127.0.0.1", 0), 서버.처리기)
    threading.Thread(target=srv.serve_forever, daemon=True).start()
    포트 = srv.server_address[1]
    try:
        def 상태(머리):
            c = http.client.HTTPConnection("127.0.0.1", 포트)
            c.request("POST", urllib.request.quote("/api/목록"), body=b"{}", headers=머리)
            return c.getresponse().status
        assert 상태({"Content-Type": "application/json"}) == 200
        assert 상태({"Content-Type": "application/json", "Host": "evil.example"}) == 403
        assert 상태({"Content-Type": "application/json", "Origin": "http://evil.example"}) == 403
        assert 상태({"Content-Type": "text/plain"}) == 415
    finally:
        srv.shutdown()


def test_v2_빈글_줄바꿈_두번쓴_원본장(tmp_path):
    import 바탕pptx
    v1 = 받은자료 / "ACME_중간보고_v1.pptx"
    아이디 = Presentation(str(v1)).slides[1].shapes[0].shape_id
    f = tmp_path / "v2.pptx"
    기록 = 바탕pptx.v2만들기(v1, [{"종류": "원본장", "원본": 1}, {"종류": "원본장", "원본": 1},
                               {"종류": "원본장", "원본": 2, "고침": {str(아이디): "첫줄\x0b둘째"}},
                               {"종류": "원본장", "원본": 9}], None, f)
    s = Presentation(str(f)).slides
    assert len(s) == 2 and "_x000B_" not in s[1]._element.xml and "<a:br/>" in s[1]._element.xml
    assert any("또 넣으려" in x for x in 기록) and any("9장은 v1에 없어" in x for x in 기록)
