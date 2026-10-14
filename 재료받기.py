"""재료 받기: 넣은 파일을 그림·글·표로 풀어 재료함에 둔다.

재료 하나 = 프로젝트/<덱>/재료/<번호_이름>/
    원본.<확장자>     넣은 파일 그대로(고치지 않는다)
    정보.json         종류, 원래 이름, 무엇을 꺼냈나
    쪽/NN.png         PDF·PPTX 쪽 그림(화면 확인용)
    그림/...          박힌 그림. 원본 형식 그대로 꺼낸다(무손실)
    글.json           글 덩어리 [{쪽, 글, 자리}]
    표/<이름>.json    표 {이름, 머리, 줄들}

PDF 회의록, PPTX(v1 등), 엑셀/CSV, 그림, mp4, Origin, 글(.md/.txt)을 받는다.
"""

from __future__ import annotations

import csv
import io
import json
import re
import shutil
import subprocess
import sys
import time
from pathlib import Path

종류표 = {
    ".pdf": "PDF", ".pptx": "PPTX", ".xlsx": "엑셀", ".xlsm": "엑셀", ".csv": "CSV",
    ".png": "그림", ".jpg": "그림", ".jpeg": "그림", ".tif": "그림", ".tiff": "그림",
    ".bmp": "그림", ".gif": "그림", ".webp": "그림",
    ".mp4": "영상", ".mov": "영상", ".avi": "영상",
    ".opju": "Origin", ".opj": "Origin",
    ".md": "글", ".txt": "글",
}


def _이름다듬기(이름: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", 이름).strip("._")[:60] or "재료"


def 넣기(프로젝트: Path, 파일: Path | None = None, *, 이름: str | None = None,
       내용: bytes | None = None) -> dict:
    """파일(경로) 또는 내용(브라우저에서 올린 것)을 재료로 넣고 풀어 둔다."""
    이름 = 이름 or 파일.name
    확장 = Path(이름).suffix.lower()
    종류 = 종류표.get(확장)
    if not 종류:
        raise ValueError(f"아직 못 받는 형식이에요: {확장 or 이름}")
    밑 = 프로젝트 / "재료"
    밑.mkdir(parents=True, exist_ok=True)
    번호 = 1 + max([int(p.name.split("_", 1)[0]) for p in 밑.iterdir()
                   if p.is_dir() and p.name.split("_", 1)[0].isdigit()] or [0])
    폴더 = 밑 / f"{번호:03d}_{_이름다듬기(Path(이름).stem)}"
    폴더.mkdir()
    원본 = 폴더 / f"원본{확장}"
    if 내용 is not None:
        원본.write_bytes(내용)
    else:
        shutil.copy2(파일, 원본)
    정보 = {"종류": 종류, "원래이름": 이름, "넣은때": time.strftime("%Y-%m-%d %H:%M"),
          "알림": []}
    try:
        {"PDF": _pdf, "PPTX": _pptx, "엑셀": _엑셀, "CSV": _csv, "그림": _그림, "영상": _영상,
         "Origin": _origin, "글": _글}[종류](원본, 폴더, 정보)
    except Exception as exc:  # 일부만 풀려도 재료는 남긴다
        정보["알림"].append(f"다 풀지 못했어요: {exc}")
    (폴더 / "정보.json").write_text(json.dumps(정보, ensure_ascii=False, indent=1), "utf-8")
    return 읽기(프로젝트, 폴더.name)


def 목록(프로젝트: Path) -> list[dict]:
    밑 = 프로젝트 / "재료"
    if not 밑.exists():
        return []
    return [읽기(프로젝트, p.name) for p in sorted(밑.iterdir()) if (p / "정보.json").exists()]


def 읽기(프로젝트: Path, 아이디: str) -> dict:
    폴더 = 프로젝트 / "재료" / 아이디
    정보 = json.loads((폴더 / "정보.json").read_text("utf-8"))
    상대 = lambda p: str(p.relative_to(프로젝트)).replace("\\", "/")  # noqa: E731
    글 = json.loads((폴더 / "글.json").read_text("utf-8")) if (폴더 / "글.json").exists() else []
    표 = [json.loads(p.read_text("utf-8")) | {"경로": 상대(p)}
          for p in sorted((폴더 / "표").glob("*.json"))] if (폴더 / "표").exists() else []
    return {"아이디": 아이디, **정보,
            "쪽": [상대(p) for p in sorted((폴더 / "쪽").glob("*.png"))],
            "그림": [상대(p) for p in sorted((폴더 / "그림").glob("*")) if p.is_file()],
            "글": 글, "표": [{"경로": t["경로"], "이름": t["이름"], "머리": t["머리"],
                           "줄수": len(t["줄들"])} for t in 표]}


# ---------------------------------------------------------------- PDF (회의록 등)
def _한글(c: str) -> bool:
    return "\uac00" <= c <= "\ud7a3"


def _줄읽기(textpage) -> list[list]:
    """PDF에 적힌 순서대로 줄을 읽는다: [위 y, 아래 y, 왼 x, 오른 x, 글]"""
    전체 = textpage.get_text_range()
    줄들, i = [], 0                      # [위 y, 아래 y, 왼 x, 오른 x, 글]
    for 줄 in re.split(r"(\r\n|\n|\r)", 전체):
        if 줄 in ("\r\n", "\n", "\r"):
            i += len(줄)
            continue
        글 = 줄.strip()
        if 글:
            처음 = i + (len(줄) - len(줄.lstrip()))
            끝 = i + len(줄.rstrip()) - 1
            try:
                l, b, _, t = textpage.get_charbox(처음)
                _, _, r, _ = textpage.get_charbox(끝)
            except Exception:
                l = b = r = t = 0
            줄들.append([t, b, l, r, 글])
        i += len(줄)
    return 줄들


def _줄간격(쪽들: list[list[list]]) -> float:
    """문서 전체에서 가장 촘촘한 쪽(하위 20%) 줄 간격 = 문단 안 줄 간격."""
    간격들 = sorted(앞[0] - 뒤[0] for 줄들 in 쪽들 for 앞, 뒤 in zip(줄들, 줄들[1:])
                 if 0 < 앞[0] - 뒤[0])
    return 간격들[len(간격들) // 5] if 간격들 else 0


def _글덩어리(줄들: list[list], 쪽: int, 보통: float) -> list[dict]:
    """줄을 문단으로 묶는다. 줄 간격이 보통보다 넉넉히 벌어지면 새 문단.
    꽉 찬 줄에서 넘어간 글은 이어 붙이고(한글 낱말이 잘린 경우 띄어 쓰지 않음),
    짧은 줄(표·목록)은 줄바꿈을 그대로 둔다."""
    if not 줄들:
        return []
    오른끝 = max(z[3] for z in 줄들)
    덩어리, 지금 = [], []
    for z in 줄들:
        if 지금:
            앞 = 지금[-1]
            간격 = 앞[0] - z[0]
            if 간격 <= 0 or 간격 > 1.25 * 보통 or abs(z[2] - 지금[0][2]) > 30:
                덩어리.append(지금)
                지금 = []
        지금.append(z)
    if 지금:
        덩어리.append(지금)
    out = []
    for d in 덩어리:
        글 = d[0][4]
        for 앞, 뒤 in zip(d, d[1:]):
            꽉참 = 앞[3] > 오른끝 - 25
            if not 꽉참:
                붙임 = "\n"            # 짧은 줄은 일부러 끊은 줄(표·목록)
            elif _한글(글[-1]) and _한글(뒤[4][0]):
                붙임 = ""              # 줄 끝에서 잘린 한글 낱말
            else:
                붙임 = " "
            글 += 붙임 + 뒤[4]
        out.append({"쪽": 쪽, "글": 글})
    return out


def _pdf(원본: Path, 폴더: Path, 정보: dict) -> None:
    import pypdfium2 as pdfium
    import pypdfium2.raw as raw
    (폴더 / "쪽").mkdir()
    (폴더 / "그림").mkdir()
    doc = pdfium.PdfDocument(str(원본))
    글들, 그림수, 작은그림, 쪽줄들 = [], 0, 0, []
    try:
        for n, page in enumerate(doc, start=1):
            page.render(scale=1600 / page.get_width()).to_pil().convert("RGB").save(
                폴더 / "쪽" / f"{n:02d}.png")
            tp = page.get_textpage()
            쪽줄들.append(_줄읽기(tp))
            tp.close()
            for k, obj in enumerate(page.get_objects(filter=(raw.FPDF_PAGEOBJ_IMAGE,)), start=1):
                w, h = obj.get_px_size()
                if w < 48 or h < 48:          # 글머리표·작은 아이콘은 건너뛴다
                    작은그림 += 1
                    continue
                # 원본 형식 그대로(JPEG는 .jpg 그대로). 안 되면 화소를 PNG로
                obj.extract(str(폴더 / "그림" / f"쪽{n:02d}_그림{k:02d}"), fb_format="png")
                그림수 += 1
        정보["쪽수"] = len(doc)
    finally:
        doc.close()
    보통 = _줄간격(쪽줄들)
    for n, 줄들 in enumerate(쪽줄들, start=1):
        글들 += _글덩어리(줄들, n, 보통)
    (폴더 / "글.json").write_text(json.dumps(글들, ensure_ascii=False, indent=1), "utf-8")
    정보["꺼낸것"] = f"{정보['쪽수']}쪽 · 글 {len(글들)}덩어리 · 그림 {그림수}장"
    if 작은그림:
        정보["알림"].append(f"아주 작은 그림 {작은그림}개(아이콘 등)는 건너뛰었어요")
    if not 글들:
        정보["알림"].append("글이 없어요. 스캔한 PDF라면 쪽 그림만 쓸 수 있어요(글자 인식은 아직 없음)")


# ---------------------------------------------------------------- PPTX (v1 등)
def _pptx(원본: Path, 폴더: Path, 정보: dict) -> None:
    from pptx import Presentation

    from pptx그림 import iter_pictures      # 앱 폴더 안에 있다(다른 폴더에 기대지 않음)
    (폴더 / "그림").mkdir()
    prs = Presentation(str(원본))
    정보["쪽수"] = len(prs.slides)          # 뒤에서 무엇이 실패해도 장 수는 남긴다
    글들, 그림수 = [], 0
    for n, 장 in enumerate(prs.slides, start=1):
        for shape in 장.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                글들.append({"쪽": n, "글": shape.text_frame.text.strip(), "자리": shape.name})
        if 장.has_notes_slide and 장.notes_slide.notes_text_frame.text.strip():
            글들.append({"쪽": n, "글": 장.notes_slide.notes_text_frame.text.strip(), "자리": "노트"})
        for k, ref in enumerate(iter_pictures(장), start=1):
            try:
                img = ref.shape.image
            except Exception:
                continue
            (폴더 / "그림" / f"장{n:02d}_그림{k:02d}.{img.ext}").write_bytes(img.blob)  # 원본 그대로
            그림수 += 1
    (폴더 / "글.json").write_text(json.dumps(글들, ensure_ascii=False, indent=1), "utf-8")
    정보["쪽수"] = len(prs.slides)
    정보["꺼낸것"] = f"{len(prs.slides)}장 · 글 {len(글들)}덩어리 · 그림 {그림수}장"
    import 굽기   # 쪽 그림: PowerPoint(있으면) 또는 LibreOffice
    try:
        쪽들, 프로그램 = 굽기.pptx그림들(원본, 폴더 / "쪽")
        for i, p in enumerate(쪽들, start=1):
            p.replace(폴더 / "쪽" / f"{i:02d}.png")
        if not 쪽들:
            정보["알림"].append("쪽 그림을 만들 프로그램(PowerPoint/LibreOffice)이 없어 글·그림만 꺼냈어요")
    except Exception as exc:
        정보["알림"].append(f"쪽 그림을 못 만들었어요: {exc}")


# ---------------------------------------------------------------- 표
def _표저장(폴더: Path, 이름: str, 줄들: list[list]) -> None:
    줄들 = [r for r in 줄들 if any(v not in (None, "") for v in r)]
    if not 줄들:
        return
    폭 = max(len(r) for r in 줄들)
    줄들 = [list(r) + [None] * (폭 - len(r)) for r in 줄들]
    머리 = [str(v) if v not in (None, "") else f"열{i + 1}" for i, v in enumerate(줄들[0])]
    (폴더 / "표").mkdir(exist_ok=True)
    (폴더 / "표" / f"{_이름다듬기(이름)}.json").write_text(
        json.dumps({"이름": 이름, "머리": 머리, "줄들": 줄들[1:]}, ensure_ascii=False, default=str),
        "utf-8")


def _엑셀(원본: Path, 폴더: Path, 정보: dict) -> None:
    import openpyxl
    wb = openpyxl.load_workbook(원본, read_only=True, data_only=True)
    for ws in wb.worksheets:
        _표저장(폴더, ws.title, [list(r) for r in ws.iter_rows(values_only=True)])
    wb.close()
    정보["꺼낸것"] = f"표 {len(list((폴더 / '표').glob('*.json')))}개"


def _csv(원본: Path, 폴더: Path, 정보: dict) -> None:
    raw = 원본.read_bytes()
    for enc in ("utf-8-sig", "cp949", "latin-1"):
        try:
            글 = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    줄들 = []
    for r in csv.reader(io.StringIO(글)):
        줄들.append([_수(v) for v in r])
    _표저장(폴더, 정보["원래이름"].rsplit(".", 1)[0], 줄들)
    정보["꺼낸것"] = "표 1개"


def _수(v: str):
    try:
        return int(v) if re.fullmatch(r"-?\d+", v.strip()) else float(v)
    except (ValueError, AttributeError):
        return v


# ---------------------------------------------------------------- 그림·글·영상·Origin
def _그림(원본: Path, 폴더: Path, 정보: dict) -> None:
    (폴더 / "그림").mkdir()
    shutil.copy2(원본, 폴더 / "그림" / Path(정보["원래이름"]).name)
    정보["꺼낸것"] = "그림 1장(원본 그대로)"


def _글(원본: Path, 폴더: Path, 정보: dict) -> None:
    글 = 원본.read_text("utf-8", errors="replace")
    덩어리 = [d.strip() for d in re.split(r"\n\s*\n", 글) if d.strip()]
    (폴더 / "글.json").write_text(json.dumps([{"쪽": 1, "글": d} for d in 덩어리],
                                           ensure_ascii=False, indent=1), "utf-8")
    정보["꺼낸것"] = f"글 {len(덩어리)}덩어리"


def ffmpeg() -> str | None:
    앱도구 = Path(__file__).parent / "도구" / ("ffmpeg.exe" if sys.platform.startswith("win") else "ffmpeg")
    return str(앱도구) if 앱도구.exists() else shutil.which("ffmpeg")


def _영상(원본: Path, 폴더: Path, 정보: dict) -> None:
    ff = ffmpeg()
    if not ff:
        정보["알림"].append("ffmpeg가 없어 영상에서 장면을 뽑지 못했어요(도구 폴더에 ffmpeg를 두면 돼요)")
        return
    (폴더 / "그림").mkdir()
    # 첫 장면 + 장면이 크게 바뀔 때마다 한 장
    subprocess.run([ff, "-hide_banner", "-loglevel", "error", "-i", str(원본), "-vf",
                    "select='eq(n,0)+gt(scene,0.3)'", "-vsync", "vfr", "-frames:v", "60",
                    str(폴더 / "그림" / "장면%03d.png")], check=True, timeout=600)
    정보["꺼낸것"] = f"장면 {len(list((폴더 / '그림').glob('*.png')))}장"


def _origin(원본: Path, 폴더: Path, 정보: dict) -> None:
    """Origin이 깔린 PC에서만: originpro로 워크시트는 표로, 그래프는 그림으로 꺼낸다."""
    try:
        import originpro as op
    except ImportError:
        정보["알림"].append("Origin 파일은 Origin이 깔린 PC에서 originpro로만 열 수 있어요. "
                          "Origin에서 데이터는 엑셀/CSV, 그래프는 PNG로 내보내 넣어도 돼요")
        return
    (폴더 / "그림").mkdir()
    try:
        op.open(str(원본), readonly=True)
        for wb in op.pages("w"):
            for ws in wb:
                df = ws.to_df()
                _표저장(폴더, f"{wb.name}_{ws.name}", [list(df.columns)] + df.values.tolist())
        for gp in op.pages("g"):
            gp.save_fig(str(폴더 / "그림" / f"{_이름다듬기(gp.name)}.png"), width=2400)
        정보["꺼낸것"] = (f"표 {len(list((폴더 / '표').glob('*.json')))}개 · "
                       f"그래프 {len(list((폴더 / '그림').glob('*.png')))}장")
    finally:
        op.exit()
