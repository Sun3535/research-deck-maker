"""바탕 PPTX(v1) 다루기: 템플릿은 그대로 두고 v2를 만든다.

- 테마읽기: v1의 테마 색·글자체 → 스타일(새 장이 v1처럼 보이게)
- 원본장들: v1 장마다 고칠 수 있는 글상자 목록
- v2만들기: v1을 열어 원본장 글 고침 + 새 장 끼워 넣기 + 순서 바꾸기 → <이름>_v2.pptx
  v1 파일은 절대 덮어쓰지 않는다. 빠진 장은 v2에만 없을 뿐 v1에 남아 있다.
"""

from __future__ import annotations

import copy
import re
from pathlib import Path

from lxml import etree
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT
from pptx.oxml.ns import qn

_A = "http://schemas.openxmlformats.org/drawingml/2006/main"


# ---------------------------------------------------------------- 테마 → 스타일
def _색값(el) -> str | None:
    if el is None:
        return None
    for c in el:
        if c.tag == f"{{{_A}}}srgbClr":
            return "#" + c.get("val").upper()
        if c.tag == f"{{{_A}}}sysClr":
            return "#" + (c.get("lastClr") or ("000000" if c.get("val") == "windowText" else "FFFFFF")).upper()
    return None


def _섞기(a: str, b: str, t: float) -> str:
    ra, rb = [int(a[i:i + 2], 16) for i in (1, 3, 5)], [int(b[i:i + 2], 16) for i in (1, 3, 5)]
    return "#" + "".join(f"{round(x + (y - x) * t):02X}" for x, y in zip(ra, rb))


def 테마읽기(pptx: Path) -> dict:
    """v1의 테마 → {"색", "계열", "글자체", "설명"} (스타일 사전과 같은 모양)."""
    prs = Presentation(str(pptx))
    master = prs.slide_master
    theme = etree.fromstring(master.part.part_related_by(RT.THEME).blob)
    cs = theme.find(f".//{{{_A}}}clrScheme")
    c = {el.tag.split("}")[1]: _색값(el) for el in cs} if cs is not None else {}
    글, 바탕 = c.get("dk1") or "#000000", c.get("lt1") or "#FFFFFF"
    강조 = c.get("accent1") or "#2A5BD7"
    fs = theme.find(f".//{{{_A}}}fontScheme")
    글자체 = None
    if fs is not None:
        minor = fs.find(f"{{{_A}}}minorFont")
        for el in ([minor.find(f"{{{_A}}}ea")] + minor.findall(f"{{{_A}}}font")) if minor is not None else []:
            if el is not None and el.get("typeface") and (el.tag.endswith("ea") or el.get("script") == "Hang"):
                글자체 = el.get("typeface")
                if el.get("script") == "Hang":
                    break
        if not 글자체 and minor is not None and minor.find(f"{{{_A}}}latin") is not None:
            글자체 = minor.find(f"{{{_A}}}latin").get("typeface")
    return {
        "설명": f"{pptx.stem} 템플릿의 테마",
        "색": {"바탕": 바탕, "면": _섞기(바탕, 글, 0.06), "글": 글, "흐린글": _섞기(글, 바탕, 0.35),
               "강조": 강조, "선": _섞기(바탕, 글, 0.18)},
        "계열": [c.get(f"accent{i}") or 강조 for i in range(1, 5)],
        "글자체": 글자체 or "맑은 고딕",
        "크기": [prs.slide_width, prs.slide_height],
    }


# ---------------------------------------------------------------- 원본장
def 원본장들(pptx: Path) -> list[dict]:
    """v1 장마다 고칠 수 있는 글상자: [{번호, 글상자: [{아이디, 이름, 글}]}]"""
    prs = Presentation(str(pptx))
    out = []
    for n, 장 in enumerate(prs.slides, start=1):
        상자 = [{"아이디": s.shape_id, "이름": s.name, "글": s.text_frame.text}
              for s in 장.shapes if s.has_text_frame and s.text_frame.text.strip()]
        out.append({"번호": n, "글상자": 상자})
    return out


def _글바꾸기(shape, 새글: str) -> None:
    """글상자의 글만 바꾸고 모양(글자체·크기·색·글머리)은 원래 것을 그대로 쓴다."""
    tf = shape.text_frame
    줄들 = 새글.split("\n")
    문단들 = list(tf.paragraphs)
    본보기 = 문단들[0]._p
    for i, 줄 in enumerate(줄들):
        if i < len(문단들):
            p = 문단들[i]
        else:   # 줄이 늘면 마지막 문단을 본떠 더한다
            새 = copy.deepcopy(문단들[-1]._p if 문단들 else 본보기)
            tf._txBody.append(새)
            p = tf.paragraphs[-1]
        # 첫 글자 모양(rPr)을 기억해 두고 글을 통째로 넣는다. p.text 는 \x0b(Shift+Enter)를 <a:br/>로 바꿔 준다
        첫 = p.runs[0]._r.find(qn("a:rPr")) if p.runs else None
        모양 = copy.deepcopy(첫) if 첫 is not None else None
        p.text = 줄
        if 모양 is not None:
            for r in p.runs:
                옛 = r._r.find(qn("a:rPr"))
                if 옛 is not None:
                    r._r.remove(옛)
                r._r.insert(0, copy.deepcopy(모양))
    for p in list(tf.paragraphs)[len(줄들):]:   # 줄이 줄면 남는 문단을 뺀다
        p._p.getparent().remove(p._p)


def 다음판이름(v1: Path) -> str:
    """발표_v1.pptx → 발표_v2.pptx, 발표.pptx → 발표_v2.pptx"""
    m = re.search(r"(.*?)[_\- ]?v(\d+)$", v1.stem, re.I)
    return f"{m.group(1)}_v{int(m.group(2)) + 1}" if m else f"{v1.stem}_v2"


def _빈판(prs):
    """새 장에 쓸 판(레이아웃): 개체 틀이 가장 적은 것(보통 '빈 화면')."""
    return min(prs.slide_layouts, key=lambda l: (len(l.placeholders), l.name != "Blank"))


def v2만들기(v1: Path, 장들: list[dict], 새장그리기, 저장경로: Path) -> list[str]:
    """장들: 설계 순서대로 {"종류": "원본장", "원본": n, "고침": {아이디: 글}} 또는 새 장 자료.

    새장그리기(prs, slide, 장자료): 새 장에 박스를 그리는 함수(그리기_pptx 가 준다).
    돌려주는 것: 변경 기록 줄들.
    """
    prs = Presentation(str(v1))
    원래 = list(prs.slides)
    기록: list[str] = []
    판 = _빈판(prs)
    순서 = []
    for i, 장 in enumerate(장들, start=1):
        if 장.get("종류") == "원본장":
            try:
                n = int(장["원본"])
            except (KeyError, TypeError, ValueError):
                n = 0
            if not 1 <= n <= len(원래):
                기록.append(f"- {i}번째 자리: 원본 {장.get('원본')}장은 v1에 없어 건너뛰었어요")
                continue
            s = 원래[n - 1]
            if s in 순서:
                기록.append(f"- {i}번째 자리: 원본 {n}장을 또 넣으려 해서 건너뛰었어요(한 장은 한 번만)")
                continue
            for 아이디, 새글 in (장.get("고침") or {}).items():
                shape = next((x for x in s.shapes if str(x.shape_id) == str(아이디)), None)
                if shape is not None and shape.has_text_frame and shape.text_frame.text != 새글:
                    _글바꾸기(shape, 새글)
                    첫줄 = (새글.splitlines() or [""])[0]
                    기록.append(f"- {i}장(원본 {n}장) 글 고침: “{첫줄[:40]}”" if 첫줄 else f"- {i}장(원본 {n}장) 글을 비움")
            if n != i:
                기록.append(f"- 원본 {n}장 → {i}번째 자리")
            순서.append(s)
        else:
            s = prs.slides.add_slide(판)
            for ph in list(s.placeholders):      # 빈 개체 틀은 지운다(새 장은 박스로 그린다)
                ph._element.getparent().remove(ph._element)
            새장그리기(prs, s, 장)
            기록.append(f"- {i}장 새로 넣음: {장.get('종류')} “{장.get('제목', '')[:40]}”")
            순서.append(s)
    빠짐 = [n for n, s in enumerate(원래, start=1) if s not in 순서]
    for n in 빠짐:
        기록.append(f"- 원본 {n}장은 v2에서 뺐어요(v1에는 그대로 있어요)")
    # 장 순서 다시 쓰기: sldIdLst 를 설계 순서로
    목록 = prs.slides._sldIdLst
    요소 = {int(el.get("id")): el for el in 목록}
    새순서 = [s.slide_id for s in 순서]                   # 목록을 비우기 전에 번호를 구해 둔다
    뺄것 = [요소[원래[n - 1].slide_id].get(qn("r:id")) for n in 빠짐]
    for el in list(목록):
        목록.remove(el)
    for i in 새순서:
        목록.append(요소[i])
    for rid in 뺄것:   # 뺀 장의 연결도 끊는다(파일 안에 고아로 남지 않게)
        prs.part.drop_rel(rid)
    저장경로.parent.mkdir(parents=True, exist_ok=True)
    prs.save(저장경로)
    return 기록
